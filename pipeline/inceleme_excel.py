"""Uzman incelemesi için Excel dosyası üretir ve doldurulmuş dosyadaki kararları soru bankasına işler.

Claude hesabı olmayan bir uzman için önizleme sayfasına alternatiftir.

    python -m pipeline.inceleme_excel olustur docs/SMMM_Pilot_Uzman_Inceleme.xlsx
    python -m pipeline.inceleme_excel isle  <doldurulmus.xlsx> [--isle]
"""

from __future__ import annotations

import argparse
import sys
from datetime import date
from pathlib import Path

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from pipeline import denetle, g2_yapi

KARARLAR = ["Doğru", "Hatalı"]
KATEGORILER = {
    "Kritik - cevap anahtarı yanlış": "kritik_anahtar",
    "Kritik - birden fazla doğru şık": "kritik_coklu",
    "Kritik - mevzuat/teknik bilgi yanlış": "kritik_mevzuat",
    "Önemli - açıklama eksik veya yanıltıcı": "onemli_aciklama",
    "Önemli - kök belirsiz": "onemli_belirsiz",
    "Önemli - sınav tarzına uymuyor": "onemli_tarz",
    "Küçük - yazım/üslup": "kucuk_yazim",
}
SUTUNLAR = [
    ("Soru No", 6), ("Kimlik", 22), ("Sürüm", 7), ("Ders", 20), ("Konu", 24), ("Zorluk", 8),
    ("Soru", 70), ("A", 28), ("B", 28), ("C", 28), ("D", 28), ("E", 28),
    ("Doğru Cevap", 9), ("Açıklama", 70), ("Dayanak", 50),
    ("KARARINIZ", 14), ("HATA TÜRÜ", 34), ("NOTUNUZ", 50),
]
UZMAN_SUTUN = {"KARARINIZ", "HATA TÜRÜ", "NOTUNUZ"}


def _aciklama(s: dict) -> str:
    a = s["aciklama"]
    parca = [f"Doğru ({s['dogru']}): {a['dogru_neden']}"]
    if a.get("hesap_adimlari"):
        parca.append("Çözüm adımları:\n" + "\n".join(f"  {i + 1}. {x}" for i, x in enumerate(a["hesap_adimlari"])))
    parca.append("Diğer şıklar:\n" + "\n".join(f"  {h}) {t}" for h, t in sorted(a["celdiriciler"].items())))
    return "\n\n".join(parca)


def olustur(hedef: Path) -> int:
    mufredat = yaml.safe_load((g2_yapi.KOK / "content/mufredat/smmm.yaml").read_text(encoding="utf-8"))
    dersler = {d["kod"]: d for d in mufredat["dersler"]}
    sorular = [s for _, s in g2_yapi.dosyalari_yukle(g2_yapi.KOK / "content/sorular/smmm")
               if s.get("durum") in ("kontrolde", "onayli")]
    sorular.sort(key=lambda s: (list(dersler).index(s["ders"]) if s["ders"] in dersler else 99, s["id"]))

    wb = Workbook()
    bilgi = wb.active
    bilgi.title = "Nasıl doldurulur"
    satirlar = [
        "SMMM Pilot Soru Seti — Uzman İnceleme Formu",
        f"Hazırlanma: {date.today():%d.%m.%Y} · Soru sayısı: {len(sorular)}",
        "",
        "1. 'Sorular' sayfasında her satır bir sorudur. Soru, şıklar, doğru cevap, açıklama ve dayanak yan yana verilmiştir.",
        "2. Sağdaki sarı sütunları doldurun:",
        "   • KARARINIZ: listeden 'Doğru' veya 'Hatalı' seçin.",
        "   • HATA TÜRÜ: yalnız 'Hatalı' ise listeden seçin.",
        "   • NOTUNUZ: 'Hatalı' ise ne yanlış ve doğrusu ne olmalı; isterseniz 'Doğru' sorulara da öneri yazabilirsiniz.",
        "3. Kimlik ve Sürüm sütunlarını değiştirmeyin; kararlar bunlarla eşleştirilir.",
        "4. Dosyayı kaydedip geri gönderin. Yarım kalan dosyayı da gönderebilirsiniz; boş satırlar incelenmemiş sayılır.",
        "",
        "Özellikle bakılmasını istediğimiz noktalar 'İnceleme rehberi' belgesindedir (TFRS 18 geçişi, yöntem soruları vb.).",
        "Bu uygulamanın TÜRMOB veya TESMER ile resmî bir bağlantısı yoktur.",
    ]
    for i, t in enumerate(satirlar, 1):
        bilgi.cell(row=i, column=1, value=t)
    bilgi["A1"].font = Font(bold=True, size=14)
    bilgi.column_dimensions["A"].width = 120

    ws = wb.create_sheet("Sorular")
    baslik_dolgu = PatternFill("solid", fgColor="0F6B5C")
    uzman_dolgu = PatternFill("solid", fgColor="FFF4C2")
    for j, (ad, gen) in enumerate(SUTUNLAR, 1):
        h = ws.cell(row=1, column=j, value=ad)
        h.font = Font(bold=True, color="FFFFFF")
        h.fill = baslik_dolgu
        h.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[h.column_letter].width = gen
    ws.freeze_panes = "H2"

    for i, s in enumerate(sorular, 2):
        d = dersler.get(s["ders"], {})
        konu = next((k["ad"] for k in d.get("konular", []) if k["kod"] == s["konu"]), s["konu"])
        dayanak = "\n".join(f"{k['mevzuat']}, {k['madde']}" for k in s["kaynaklar"])
        degerler = [i - 1, s["id"], s["surum"], d.get("ad", s["ders"]), konu, s["zorluk"],
                    s["kok"].replace("**", ""), *[s["secenekler"][h] for h in "ABCDE"],
                    s["dogru"], _aciklama(s), dayanak, None, None, None]
        for j, v in enumerate(degerler, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if SUTUNLAR[j - 1][0] in UZMAN_SUTUN:
                c.fill = uzman_dolgu

    son = len(sorular) + 1
    karar_dv = DataValidation(type="list", formula1='"' + ",".join(KARARLAR) + '"', allow_blank=True)
    ws.add_data_validation(karar_dv)
    karar_dv.add(f"P2:P{son}")
    liste = wb.create_sheet("Listeler")
    for i, k in enumerate(KATEGORILER, 1):
        liste.cell(row=i, column=1, value=k)
    liste.sheet_state = "hidden"
    kat_dv = DataValidation(type="list", formula1=f"=Listeler!$A$1:$A${len(KATEGORILER)}", allow_blank=True)
    ws.add_data_validation(kat_dv)
    kat_dv.add(f"Q2:Q{son}")

    hedef.parent.mkdir(parents=True, exist_ok=True)
    wb.save(hedef)
    return len(sorular)


def isle(kaynak: Path, yaz: bool) -> int:
    ws = load_workbook(kaynak, data_only=True)["Sorular"]
    basliklar = [c.value for c in ws[1]]
    sutun = {ad: basliklar.index(ad) for ad in ("Kimlik", "Sürüm", "KARARINIZ", "HATA TÜRÜ", "NOTUNUZ")}
    kayitlar = g2_yapi.dosyalari_yukle(g2_yapi.KOK / "content/sorular/smmm", [])
    sorular = {s["id"]: s for _, s in kayitlar}
    sayac = {"dogru": 0, "hatali": 0, "bos": 0, "eski": 0, "eksik_tur": 0}
    islenen = set()
    for satir in ws.iter_rows(min_row=2, values_only=True):
        sid, surum = satir[sutun["Kimlik"]], satir[sutun["Sürüm"]]
        karar = (satir[sutun["KARARINIZ"]] or "").strip()
        tur = (satir[sutun["HATA TÜRÜ"]] or "").strip()
        not_ = (satir[sutun["NOTUNUZ"]] or "").strip()
        soru = sorular.get(sid)
        if not soru:
            continue
        if not karar:
            sayac["bos"] += 1
            continue
        if surum != soru["surum"]:
            sayac["eski"] += 1
            print(f"  ESKİ SÜRÜM {sid}: formdaki {surum}, bankadaki {soru['surum']}")
            continue
        if karar == "Hatalı":
            sayac["hatali"] += 1
            if tur not in KATEGORILER:
                sayac["eksik_tur"] += 1
            print(f"  HATALI {sid} [{tur or 'tür seçilmemiş'}] {not_}")
            kayit = {"sonuc": "kaldi", "tarih": str(date.today()), "not": f"{KATEGORILER.get(tur, tur)}: {not_}"}
        else:
            sayac["dogru"] += 1
            kayit = {"sonuc": "gecti", "tarih": str(date.today()), **({"not": "Öneri: " + not_} if not_ else {})}
        if yaz:
            soru.setdefault("uretim", {}).setdefault("kapilar", {})["G8"] = kayit
            if karar == "Doğru" and soru.get("durum") == "kontrolde":
                soru["durum"] = "onayli"
            islenen.add(sid)
    if yaz:
        dosyalar = {d for d, s in kayitlar if s["id"] in islenen}
        denetle._yaz([(d, s) for d, s in kayitlar if d in dosyalar], g2_yapi.KOK / "content/sorular/smmm")
    print(f"Doğru: {sayac['dogru']} · Hatalı: {sayac['hatali']} · İncelenmemiş: {sayac['bos']} · "
          f"Eski sürüm: {sayac['eski']}" + (f" · Hata türü seçilmemiş: {sayac['eksik_tur']}" if sayac["eksik_tur"] else ""))
    return 0


def main(argv: list[str]) -> int:
    ap = argparse.ArgumentParser()
    alt = ap.add_subparsers(dest="komut", required=True)
    o = alt.add_parser("olustur")
    o.add_argument("hedef", type=Path)
    i = alt.add_parser("isle")
    i.add_argument("kaynak", type=Path)
    i.add_argument("--isle", action="store_true")
    args = ap.parse_args(argv[1:])
    if args.komut == "olustur":
        n = olustur(args.hedef)
        print(f"{n} soru → {args.hedef}")
        return 0
    return isle(args.kaynak, args.isle)


if __name__ == "__main__":
    sys.exit(main(sys.argv))
